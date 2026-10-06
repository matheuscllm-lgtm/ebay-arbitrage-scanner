"""Regressões offline de T7: UTF-8, caminhos externos e dados literais no XLSX.

Extrai funções reais via AST para não executar coletas nem carregar pickle.
Todos os arquivos gerados são fixtures temporárias.
"""
import ast
import builtins
import json
import os
import re
import sys
import tempfile
import unittest
import urllib.parse
import zipfile
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent / 'pipeline'
sys.path.insert(0, str(ROOT))
import common


def functions(filename, names, namespace):
    tree = ast.parse((ROOT / filename).read_text(encoding='utf-8'))
    nodes = [n for n in tree.body if isinstance(n, ast.FunctionDef) and n.name in names]
    exec(compile(ast.Module(body=nodes, type_ignores=[]), filename, 'exec'), namespace)
    return namespace


class Utf8Tests(unittest.TestCase):
    def test_load_preserves_unicode_with_windows_default(self):
        original_open = builtins.open
        def windows_open(file, mode='r', *args, **kwargs):
            if 'b' not in mode and 'encoding' not in kwargs:
                kwargs['encoding'] = 'cp1252'
            return original_open(file, mode, *args, **kwargs)
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            name = 'Pokémon 皮卡丘 Nidoran♀'
            for filename, data in [('collection_manifest.json', {'complete': True}),
                                   ('sets.json', [{'id': 1, 'name': name}]),
                                   ('all_cards.json', [{'name': name}])]:
                (root / filename).write_text(json.dumps(data, ensure_ascii=False), encoding='utf-8')
            cwd = os.getcwd()
            try:
                os.chdir(root)
                with patch('builtins.open', windows_open):
                    sets, cards = common.load()
                self.assertEqual(sets[1]['name'], name)
                self.assertEqual(cards[0]['name'], name)
            finally:
                os.chdir(cwd)

    def test_pipeline_text_io_declares_encoding(self):
        # fetch_cards.py pertence a T2 e foi explicitamente excluído de T7.
        for path in ROOT.glob('*.py'):
            if path.name == 'fetch_cards.py':
                continue
            for call in ast.walk(ast.parse(path.read_text(encoding='utf-8'))):
                if not isinstance(call, ast.Call):
                    continue
                mode = 'r'
                if isinstance(call.func, ast.Name) and call.func.id == 'open':
                    value = next((k.value for k in call.keywords if k.arg == 'mode'),
                                 call.args[1] if len(call.args) > 1 else ast.Constant('r'))
                    self.assertIsInstance(value, ast.Constant)
                    mode = value.value
                elif not (isinstance(call.func, ast.Attribute) and call.func.attr in ('read_text', 'write_text')):
                    continue
                if 'b' in mode:
                    continue
                with self.subTest(path=path.name, line=call.lineno):
                    encoding = next((k.value.value for k in call.keywords if k.arg == 'encoding'), None)
                    self.assertIn(encoding, ('utf-8', 'utf-8-sig'))


class ExternalPathTests(unittest.TestCase):
    def setUp(self):
        self.scope = functions('fetch_external.py', {'validated_set_id', 'ptcg_job', 'tcgdex_job'},
                               {'re': re, 'urllib': urllib, 'RAW': 'https://fixture.invalid/data'})

    def test_valid_ids_preserved_and_url_encoded(self):
        for sid in ('base1', 'SV-P', 'SM4+', 'CSV9.5C', 'sv4pt5', 'set_1'):
            with self.subTest(sid=sid):
                url, path = self.scope['ptcg_job']({'id': sid})
                self.assertEqual(path, f'ext/ptcg/{sid}.json')
                self.assertIn(urllib.parse.quote(sid, safe=''), url)
                url, path = self.scope['tcgdex_job']('zh-tw', sid)
                self.assertEqual(path, f'ext/tcgdex_sets/zh-tw__{sid}.json')
                self.assertIn(urllib.parse.quote(sid, safe=''), url)

    def test_traversal_absolute_and_windows_paths_rejected(self):
        invalid = ('../outside', '/tmp/out', 'x/y', r'x\y', r'C:\out', 'a/../../b',
                   '..', 'a..b', 'a%2fb', 'a%5cb', 'a%2e%2e', 'a?b', 'a#b', 'a\x00b',
                   '', None, 1, ' x', 'x ', 'CON', 'con.txt', 'LPT1', 'NUL', 'end.', 'x' * 101)
        for sid in invalid:
            with self.subTest(sid=sid):
                with self.assertRaises(ValueError):
                    self.scope['ptcg_job']({'id': sid})
                with self.assertRaises(ValueError):
                    self.scope['tcgdex_job']('ja', sid)
        with self.assertRaises(ValueError):
            self.scope['tcgdex_job']('../en', 'base1')


class SpreadsheetLiteralTests(unittest.TestCase):
    def test_saved_xlsx_keeps_data_literal_and_internal_formulas_active(self):
        wb = Workbook()
        scope = functions('build_xlsx.py', {'literal_cell', 'sheet'},
                          {'wb': wb, 'HF': Font(bold=True), 'HFILL': PatternFill(),
                           'Alignment': Alignment, 'get_column_letter': get_column_letter})
        values = ['=HYPERLINK("https://fixture.invalid", "x")', '+SUM(1,2)',
                  '-1+2', '@SUM(1,2)', 'Pokémon 皮卡丘', -2, 0, None]
        ws = scope['sheet']('Dados', ['Valor'], [[v] for v in values])
        # Era recebida no Leia-me passa pelo mesmo escape.
        scope['literal_cell'](wb.active, 1, 1, '=1+2')
        ws['B2'] = '=SUM(1,2)'  # fórmula interna deliberada
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'fixture.xlsx'
            wb.save(path)
            with zipfile.ZipFile(path) as archive:
                xml = archive.read('xl/worksheets/sheet2.xml').decode('utf-8')
            self.assertEqual(xml.count('<f>'), 1)
            loaded = load_workbook(path, data_only=False)
            try:
                for index, value in enumerate(values, 2):
                    cell = loaded['Dados'].cell(index, 1)
                    self.assertEqual(cell.value, value)
                    self.assertNotEqual(cell.data_type, 'f')
                    if index < 6:
                        self.assertTrue(cell.quotePrefix)
                self.assertEqual(loaded['Dados']['B2'].data_type, 'f')
                self.assertEqual(loaded.active['A1'].data_type, 's')
            finally:
                loaded.close()
            loaded = load_workbook(path, data_only=True)
            try:
                self.assertEqual(loaded['Dados']['A2'].value, values[0])
            finally:
                loaded.close()


if __name__ == '__main__':
    unittest.main()
