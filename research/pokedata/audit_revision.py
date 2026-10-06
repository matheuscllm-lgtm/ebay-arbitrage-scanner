"""Read-only audit of the revised private workbook against the original PR53 file.

Usage: python audit_revision.py revised.xlsx original.xlsx comparison.csv
Does not load pickle files, download sources, modify workbooks or print prices.
"""
import csv
from collections import Counter
import json
import sys
from openpyxl import load_workbook

def audit(revised, original, comparison):
    new = load_workbook(revised, read_only=True, data_only=True)
    old = load_workbook(original, read_only=True, data_only=True)
    try:
        coverage = list(new['Cobertura PR53'].values)[1:]
        original_refs = [r for r in old['Cobertura EN'].iter_rows(min_row=5, values_only=True) if r[0] is not None]
        assert len(coverage) == 3490 == len({r[0] for r in coverage})
        assert Counter(tuple(r[:4]) for r in coverage) == Counter(tuple(r[:4]) for r in original_refs)
        counts = Counter(r[7] for r in coverage)
        assert counts == {'arte confirmada': 2516, 'provável': 323, 'inconclusivo': 479, 'não encontrada': 172}
        old_pairs = [r for r in old['Detalhes confirmados'].iter_rows(min_row=5, values_only=True) if r[0] is not None]
        cht = [r for r in old_pairs if r[5] == 'Chinês tradicional']
        revised_cht = list(new['CHT PR53'].values)[1:]
        assert len(cht) == len(revised_cht) == 116
        assert Counter(tuple(r[i] for i in [0,1,2,3,6,7,8,9,10,11,12]) for r in cht) == Counter(tuple(r[:11]) for r in revised_cht)
        with open(comparison, encoding='utf-8-sig', newline='') as stream:
            compared = list(csv.DictReader(stream))
        langs = {'Japonês': 'JP', 'Chinês simplificado': 'CHS', 'Chinês tradicional': 'CHT'}
        assert len(compared) == len(old_pairs) == 388
        assert Counter((int(r['id_en']),r['idioma'],r['codigo_pr53']) for r in compared) == Counter((r[0],langs[r[5]],r[7]) for r in old_pairs)
        cross = list(new['Correspondência'].values)[1:]
        assert len(cross) == 18345 and sum(r[34] for r in cross) == 18673
        jp = sum(r[7] == 'arte confirmada' for r in cross)
        chs = sum(r[16] == 'arte confirmada' for r in cross)
        both = sum(r[7] == r[16] == 'arte confirmada' for r in cross)
        assert (jp,chs,both) == (17707,7189,6551)
        assert jp + chs - both == len(cross)
        return {'references':len(coverage),'coverage_status':dict(counts),
                'comparison_records':len(compared),'unique_reference_language_pairs':len({(int(r['id_en']),r['idioma']) for r in compared}),
                'cht_fields_preserved':True,'correspondence_rows':len(cross),
                'units_in_rows':sum(r[34] for r in cross),
                'confirmed_reference_image':dict(Counter(r[6] for r in coverage if r[7]=='arte confirmada')),
                'limitation':'Counts and preservation only; no image revalidation or formula recalculation.'}
    finally:
        new.close();old.close()

if __name__ == '__main__':
    if len(sys.argv) != 4:
        raise SystemExit(__doc__)
    print(json.dumps(audit(*sys.argv[1:]), ensure_ascii=False, indent=2))
