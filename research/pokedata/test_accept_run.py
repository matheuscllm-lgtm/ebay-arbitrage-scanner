"""Testes offline das conferências de aceite de uma execução nova (accept_run.py)."""
import unittest
import csv
import io
import json
import tempfile
from contextlib import redirect_stdout
from pathlib import Path

from openpyxl import Workbook

from accept_run import (check_comparison, check_correspondence, check_coverage, check_cht,
                        check_exclusives, check_manifest, check_readme_counts, main, CHT_FIELDS)


def ref(i, nome='Pikachu', cod='25/102', st='Base Set'):
    return {'ID EN': i, 'Carta inglesa — referência': nome, 'Código EN': cod, 'Set EN': st}


def cov(i, loc='único', geral='arte confirmada', **kw):
    row = {'ID EN (PR #53)': i, 'Carta inglesa — referência': 'Pikachu', 'Código EN': '25/102', 'Set EN': 'Base Set',
           'Localização neste catálogo': loc, 'Situação geral (esta rodada)': geral}
    row.update(kw)
    return row


class ManifestTests(unittest.TestCase):
    def test_complete_without_failures(self):
        ok, info = check_manifest({'complete': True, 'sets': [{'status': 'cached'}, {'status': 'downloaded'}], 'cards': 2})
        self.assertTrue(ok); self.assertEqual(info['sets'], 2)

    def test_incomplete_or_failed(self):
        self.assertFalse(check_manifest({'complete': False, 'sets': []})[0])
        self.assertFalse(check_manifest({'complete': True, 'sets': [{'status': 'failed'}]})[0])
        self.assertFalse(check_manifest(None)[0])


class ExclusiveTests(unittest.TestCase):
    def test_no_automatic_exclusive(self):
        catalogs = {'EN': [{'Carta única': 1, 'Situação geral': 'arte confirmada'}]}
        ok, info = check_exclusives(catalogs, [])
        self.assertTrue(ok); self.assertEqual(info['exclusivas'], 0)

    def test_exclusive_flagged(self):
        catalogs = {'JP': [{'Carta única': 1, 'Situação geral': 'exclusiva'}]}
        ok, info = check_exclusives(catalogs, [{'Critério': 'x'}])
        self.assertFalse(ok); self.assertEqual(info['exclusivas'], 1)

    def test_exclusive_on_non_primary_variant_fails(self):
        self.assertFalse(check_exclusives({'EN': [{'Carta única': 0, 'Situação geral': 'exclusiva'}]}, [])[0])


class CoverageTests(unittest.TestCase):
    def test_same_references_and_counts(self):
        original = [ref(1), ref(2, 'Raichu', '26/102')]
        rows = [cov(1), cov(2, loc='ambíguo: 2 registros', geral='')]
        rows[1].update({'Carta inglesa — referência': 'Raichu', 'Código EN': '26/102'})
        ok, info = check_coverage(rows, original)
        self.assertTrue(ok)
        self.assertEqual(info['referencias'], 2)
        self.assertEqual(info['ambiguas'], [2]); self.assertEqual(info['nao_localizadas'], [])
        self.assertEqual(info['situacao_geral'], {'arte confirmada': 1, '': 1})

    def test_reference_mismatch_fails(self):
        ok, info = check_coverage([cov(1)], [ref(1, nome='Outra')])
        self.assertFalse(ok)

    def test_duplicate_id_fails(self):
        ok, _ = check_coverage([cov(1), cov(1)], [ref(1), ref(1)])
        self.assertFalse(ok)


class ChtTests(unittest.TestCase):
    def test_preserved_fields(self):
        pair = {'ID EN': 1, 'Carta inglesa — referência': 'Pikachu', 'Código EN': '25/102', 'Set EN': 'Base Set',
                'Idioma': 'Chinês tradicional', 'Nome local': '皮卡丘', 'Código local completo': 'CS1a 1', 'Set / produto local': 'CS1a',
                'Raridade e acabamento': 'C', 'Status': 'Confirmada', 'Como foi validada': 'foto', 'Fonte local e complemento': 'loja'}
        row = {'ID EN (PR #53)': 1, 'Carta inglesa — referência': 'Pikachu', 'Código EN': '25/102', 'Set EN': 'Base Set',
               'Nome local (繁體)': '皮卡丘', 'Código local completo': 'CS1a 1', 'Set / produto local': 'CS1a', 'Raridade e acabamento': 'C',
               'Status (PR #53)': 'Confirmada', 'Como foi validada (PR #53)': 'foto', 'Fonte local': 'loja'}
        ok, info = check_cht([row], [pair, dict(pair, Idioma='Japonês')])
        self.assertTrue(ok); self.assertEqual(info['cht'], 1)
        self.assertFalse(check_cht([dict(row, **{'Nome local (繁體)': 'x'})], [pair])[0])


class ReadmeTests(unittest.TestCase):
    def setUp(self):
        self.rows = [cov(1, loc='ambíguo: 2 registros'), cov(2, loc='não localizado'), cov(3)]
        self.header = list(self.rows[0])
        self.labels = ['Não localizadas neste catálogo', 'Ambíguas: mais de um registro possível, nenhum escolhido']

    def test_correct_numeric_counts(self):
        ok, info = check_readme_counts([[label, 1] for label in self.labels], self.rows, self.header)
        self.assertTrue(ok)
        self.assertEqual(info['ambíguo']['esperado'], 1)

    def test_missing_wrong_or_duplicate_counts_fail(self):
        for readme in ([], [[label, 0] for label in self.labels],
                       [[label, 1] for label in self.labels] + [[self.labels[0], 1]]):
            self.assertFalse(check_readme_counts(readme, self.rows, self.header)[0])

    def test_generated_formulas_checked_without_recalculation(self):
        readme = [[label, f'=COUNTIF(\'Cobertura PR53\'!$E:$E,"{prefix}*")']
                  for label, prefix in zip(self.labels, ['não localizado', 'ambíguo'])]
        ok, info = check_readme_counts(readme, self.rows, self.header)
        self.assertTrue(ok)
        self.assertEqual(info['ambíguo']['modo'], 'fórmula sem recálculo')
        readme[0][1] = readme[0][1].replace('$E', '$F')
        self.assertFalse(check_readme_counts(readme, self.rows, self.header)[0])


class ComparisonTests(unittest.TestCase):
    def test_same_pairs_and_summary(self):
        pairs = [{'ID EN': 1, 'Idioma': 'Japonês', 'Código local completo': 'SV1 1'},
                 {'ID EN': 1, 'Idioma': 'Chinês simplificado', 'Código local completo': 'CS1 1'}]
        compared = [{'id_en': '1', 'idioma': 'JP', 'codigo_pr53': 'SV1 1', 'resultado': 'igual (arte confirmada aqui)'},
                    {'id_en': '1', 'idioma': 'CHS', 'codigo_pr53': 'CS1 1', 'resultado': 'sem confirmação aqui: x'}]
        ok, info = check_comparison(compared, pairs)
        self.assertTrue(ok)
        self.assertEqual(info['registros'], 2)
        self.assertEqual(info['resultado']['JP | igual (arte confirmada aqui)'], 1)
        self.assertFalse(check_comparison(compared[:1], pairs)[0])


class CorrespondenceTests(unittest.TestCase):
    def test_identity_holds(self):
        cross = [{'Situação JP': 'arte confirmada', 'Situação CN simplificado': '', 'Cartas do PokeData nesta linha': 2},
                 {'Situação JP': 'arte confirmada', 'Situação CN simplificado': 'arte confirmada', 'Cartas do PokeData nesta linha': 1},
                 {'Situação JP': '', 'Situação CN simplificado': 'arte confirmada', 'Cartas do PokeData nesta linha': 1}]
        ok, info = check_correspondence(cross)
        self.assertTrue(ok)
        self.assertEqual((info['linhas'], info['jp'], info['chs'], info['ambos'], info['cartas']), (3, 2, 2, 1, 4))

    def test_row_without_confirmed_side_fails(self):
        ok, _ = check_correspondence([{'Situação JP': 'provável', 'Situação CN simplificado': '', 'Cartas do PokeData nesta linha': 1}])
        self.assertFalse(ok)


class MainTests(unittest.TestCase):
    def test_workbook_acceptance_and_missing_readme_failure(self):
        """CLI real em fixtures mínimas: 7 checks, sem snapshots privados nem recálculo."""
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            (root / 'collection_manifest.json').write_text(
                json.dumps({'complete': True, 'sets': [{'status': 'downloaded'}], 'cards': 1}), encoding='utf-8')
            old = Workbook()
            old.remove(old.active)
            ws = old.create_sheet('Cobertura EN')
            for _ in range(3):
                ws.append(['fixture'])
            ws.append(list(ref(1)))
            ws.append(list(ref(1).values()))
            pair = {'ID EN': 1, 'Idioma': 'Chinês tradicional', **{a: 'fixture' for a, _ in CHT_FIELDS}}
            pair['Nome local'] = '皮卡丘'
            pair['Código local completo'] = 'CS1a 1'
            ws = old.create_sheet('Detalhes confirmados')
            for _ in range(3):
                ws.append(['fixture'])
            ws.append(list(pair))
            ws.append(list(pair.values()))
            old.save(root / 'old.xlsx')
            new = Workbook()
            new.remove(new.active)
            for lang in ('EN', 'JP', 'CN'):
                ws = new.create_sheet('Catálogo ' + lang)
                ws.append(['Carta única', 'Situação geral'])
                ws.append([1, 'arte confirmada'])
            new.create_sheet('Exclusivas').append(['Critério'])
            ws = new.create_sheet('Cobertura PR53')
            ws.append(list(cov(1)))
            ws.append(list(cov(1).values()))
            ws = new.create_sheet('CHT PR53')
            ws.append(['ID EN (PR #53)'] + [b for _, b in CHT_FIELDS])
            ws.append([1] + [pair[a] for a, _ in CHT_FIELDS])
            ws = new.create_sheet('Correspondência')
            ws.append(['Situação JP', 'Situação CN simplificado', 'Cartas do PokeData nesta linha'])
            ws.append(['arte confirmada', '', 1])
            readme = new.create_sheet('Leia-me')
            for label, prefix in [('Não localizadas neste catálogo', 'não localizado'),
                                  ('Ambíguas: mais de um registro possível, nenhum escolhido', 'ambíguo')]:
                readme.append([label, f'=COUNTIF(\'Cobertura PR53\'!$E:$E,"{prefix}*")'])
            new.save(root / 'new.xlsx')
            with (root / 'comp.csv').open('w', encoding='utf-8-sig', newline='') as stream:
                writer = csv.DictWriter(stream, fieldnames=['id_en', 'idioma', 'codigo_pr53', 'resultado'])
                writer.writeheader()
                writer.writerow({'id_en': 1, 'idioma': 'CHT', 'codigo_pr53': 'CS1a 1', 'resultado': 'preservado'})
            args = (root, root / 'new.xlsx', root / 'old.xlsx', root / 'comp.csv')
            output = io.StringIO()
            with redirect_stdout(output):
                self.assertEqual(main(*args), 0)
            report = json.loads(output.getvalue())
            self.assertEqual(len([r for r in report.values() if 'ok' in r]), 7)
            self.assertTrue(report['cht_pr53']['ok'])
            readme.delete_rows(2)
            new.save(root / 'new.xlsx')
            output = io.StringIO()
            with redirect_stdout(output):
                self.assertEqual(main(*args), 1)
            self.assertFalse(json.loads(output.getvalue())['leia_me_localizacao']['ok'])


if __name__ == '__main__':
    unittest.main()
