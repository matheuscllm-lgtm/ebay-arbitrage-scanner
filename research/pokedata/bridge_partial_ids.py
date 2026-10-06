"""Ponte offline entre a base parcial (3.490 referências EN) e o catálogo completo.

`ID EN` da base parcial é o número da linha (1..3490), não o `ID PokeData` do
catálogo: dos 1.785 inteiros que coincidem, só 2 têm o mesmo set e número.
Esta ponte usa metadados (set + número impresso + nome PokeData) e nunca a
igualdade de inteiros. Toda referência sai no resultado, inclusive sem match.

Uso: python research/pokedata/bridge_partial_ids.py [-o ponte.csv]
O CSV contém só identidade (IDs, set, número, nome); nunca preços.
Status: `unico` = exatamente 1 ID; `unico_numero_literal` = o catálogo cadastra a
mesma carta duas vezes ('4' e '004') e a grafia exata do número da referência escolhe
um dos registros (os dois ficam em `cadastro_duplo`); `aproximado` = só o nome-base
casa (variante não conferida); `ambiguo` = 2+ candidatos sem desempate; `sem_match`
= nenhum. As regras ficam em `pipeline/reference_match.py`, as mesmas do pipeline
(`build_xlsx.py`, `compare_pr53.py`). Nada aqui confirma arte ou impressão.
"""
import argparse
import csv
import io
import sys
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / "pipeline"))
from common import key, split_name  # noqa: E402
from reference_match import number_norm, resolve_reference  # noqa: E402

CATALOG_MEMBER = "pokedata_crossref/output/PokeData_catalogo_correspondencia.xlsx"
PARTIAL = ROOT / "inputs/PokeData_correspondencias_JP_CHS_CHT_parcial_recebido.xlsx"
FIELDS = ["id_en_parcial", "carta_parcial", "codigo_parcial", "set_parcial",
          "metodo", "status", "ids_pokedata", "numeros_impressos_catalogo", "nomes_catalogo",
          "situacao", "cadastro_duplo"]
STATUS = {"registro exato": "unico", "registro exato (número literal)": "unico_numero_literal",
          "mesma carta, registro aproximado": "aproximado"}


def text_key(value):
    """Chave de nome do pipeline (`common.key`): ♀/♂ preservados como 'f'/'m'."""
    return key(str(value or ""))


def number_key(value):
    """Número comparável do pipeline (`reference_match.number_norm`)."""
    return number_norm(value)


def build_bridge(references, catalog):
    """references: dicts com ID EN/Carta/Código EN/Set EN; catalog: linhas do Catálogo EN."""
    index = defaultdict(list)
    for row in catalog:
        index[(text_key(row["Set"]), number_key(row["Número"]))].append(dict(
            id=row["ID PokeData"], nome=row["Nome (PokeData)"], numero=row["Número"], row=row,
            unidade=(row["set_id"], str(row["Número"]), text_key(row["Nome base"]))))
    bridge = []
    for ref in references:
        cands = index.get((text_key(ref["Set EN"]), number_key(ref["Código EN"])), [])
        unit, hit, how = resolve_reference(ref["Carta inglesa — referência"], ref["Código EN"], cands)
        status = STATUS.get(how, "ambiguo" if how.startswith("ambíguo") else "sem_match")
        same_name = [c for c in cands if text_key(c["nome"]) == text_key(ref["Carta inglesa — referência"])]
        if hit is not None:
            hits = [hit]
        elif unit is not None:
            hits = [c for c in cands if c["unidade"] == unit]
        elif status == "ambiguo":  # candidatos listados para revisão; nenhum é a identidade
            base = text_key(split_name(str(ref["Carta inglesa — referência"]))[0])
            hits = same_name or [c for c in cands if c["unidade"][2] == base]
        else:
            hits = []
        bridge.append({
            "id_en_parcial": ref["ID EN"], "carta_parcial": ref["Carta inglesa — referência"],
            "codigo_parcial": ref["Código EN"], "set_parcial": ref["Set EN"],
            "metodo": "set+numero+nome_pokedata", "status": status,
            "ids_pokedata": ";".join(str(i) for i in sorted({h["id"] for h in hits})),
            "numeros_impressos_catalogo": ";".join(sorted({str(h["row"]["Número impresso"]) for h in hits})),
            "nomes_catalogo": ";".join(sorted({str(h["nome"]) for h in hits})),
            "situacao": how,
            "cadastro_duplo": (";".join(str(i) for i in sorted({c["id"] for c in same_name}))
                               if status == "unico_numero_literal" else ""),
        })
    return bridge


def id_namespace_check(references, catalog):
    """Quantos inteiros coincidem e quantos desses apontam para a mesma carta."""
    by_id = {row["ID PokeData"]: row for row in catalog}
    overlap = [ref for ref in references if ref["ID EN"] in by_id]
    same = [ref for ref in overlap
            if text_key(by_id[ref["ID EN"]]["Set"]) == text_key(ref["Set EN"])
            and number_key(by_id[ref["ID EN"]]["Número"]) == number_key(ref["Código EN"])]
    return {"numeric_overlap": len(overlap), "overlap_same_set_and_number": len(same)}


def _records(workbook, sheet, header_row):
    rows = workbook[sheet].iter_rows(min_row=header_row, values_only=True)
    header = next(rows)
    return [dict(zip(header, row)) for row in rows if any(v is not None for v in row)]


def load_inputs():
    from openpyxl import load_workbook
    with zipfile.ZipFile(ROOT / "inputs/pokedata_crossref.zip") as archive:
        data = archive.read(CATALOG_MEMBER)
    w = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    catalog = _records(w, "Catálogo EN", 1)
    w.close()
    w = load_workbook(PARTIAL, read_only=True, data_only=True)
    references = _records(w, "Cobertura EN", 4)
    w.close()
    return references, catalog


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("-o", "--output", help="CSV de saída (gitignored); omitido = só resumo")
    args = parser.parse_args(argv)
    references, catalog = load_inputs()
    bridge = build_bridge(references, catalog)
    if args.output:
        with open(args.output, "w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=FIELDS)
            writer.writeheader()
            writer.writerows(bridge)
    summary = dict(Counter(row["status"] for row in bridge))
    print({"references": len(bridge), "status": summary,
           "id_namespace": id_namespace_check(references, catalog)})
    return 0


if __name__ == "__main__":
    sys.exit(main())
