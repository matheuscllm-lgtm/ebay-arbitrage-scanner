"""Conferências de aceite de uma execução nova do pipeline (docs/POKEDATA_REPROCESSAMENTO.md, §3).

Uso: python accept_run.py <pasta de trabalho> <catalogo.xlsx> <planilha do PR #53> <comparacao_pr53.csv>

Sem contagens fixas de entregas anteriores (ao contrário de audit_revision.py): verifica
invariantes e imprime os agregados para registro. Não lê pickle, não baixa nada, não imprime preços.
Código de saída 1 quando alguma conferência falha.
"""
import csv
import json
import sys
from collections import Counter
from pathlib import Path
from openpyxl.utils import get_column_letter

REF_FIELDS = ('Carta inglesa — referência', 'Código EN', 'Set EN')
# (campo na base parcial, campo na aba CHT PR53)
CHT_FIELDS = (('Carta inglesa — referência', 'Carta inglesa — referência'), ('Código EN', 'Código EN'), ('Set EN', 'Set EN'),
              ('Nome local', 'Nome local (繁體)'), ('Código local completo', 'Código local completo'),
              ('Set / produto local', 'Set / produto local'), ('Raridade e acabamento', 'Raridade e acabamento'),
              ('Status', 'Status (PR #53)'), ('Como foi validada', 'Como foi validada (PR #53)'),
              ('Fonte local e complemento', 'Fonte local'))
LANGS = {'Japonês': 'JP', 'Chinês simplificado': 'CHS', 'Chinês tradicional': 'CHT'}


def check_manifest(manifest):
    """fetch_cards.py recusa agregado parcial; aqui só se confirma que a coleta terminou completa."""
    if not manifest:
        return False, {'erro': 'collection_manifest.json ausente'}
    failed = [r for r in manifest.get('sets', []) if r.get('status') == 'failed']
    info = {'complete': manifest.get('complete'), 'sets': len(manifest.get('sets', [])),
            'falhos': len(failed), 'cartas': manifest.get('cards')}
    return manifest.get('complete') is True and not failed, info


def check_exclusives(catalogs, exclusives):
    """Ausência não prova exclusividade (identity_policy.overall_status): nenhum status 'exclusiva' automático."""
    auto = sum(r.get('Situação geral') == 'exclusiva'
               for rows in catalogs.values() for r in rows)
    info = {'exclusivas': auto, 'aba_exclusivas': len(exclusives),
            'criterios': dict(Counter(r.get('Critério') for r in exclusives))}
    return auto == 0 and not exclusives, info


def check_coverage(rows, original_refs):
    """Cobertura PR53: as mesmas referências (ID, nome, número, set) da base parcial, uma linha por ID."""
    ids = [r['ID EN (PR #53)'] for r in rows]
    mine = Counter((r['ID EN (PR #53)'],) + tuple(r[f] for f in REF_FIELDS) for r in rows)
    theirs = Counter((r['ID EN'],) + tuple(r[f] for f in REF_FIELDS) for r in original_refs)

    def loc(r):
        return str(r.get('Localização neste catálogo') or '')
    info = {'referencias': len(rows), 'base_parcial': len(original_refs), 'ids_unicos': len(set(ids)),
            'situacao_geral': dict(Counter(r.get('Situação geral (esta rodada)') or '' for r in rows)),
            'localizacao': dict(Counter(loc(r).split(':')[0] for r in rows)),
            'ambiguas': [r['ID EN (PR #53)'] for r in rows if loc(r).startswith('ambíguo')],
            'nao_localizadas': [r['ID EN (PR #53)'] for r in rows if loc(r).startswith('não localizado')]}
    ok = len(rows) == len(original_refs) == len(set(ids)) and mine == theirs
    return ok, info


def check_readme_counts(readme_rows, coverage_rows, coverage_header):
    """§3: ambíguas/não localizadas constam no Leia-me, por valor ou fórmula correta.

    Fórmulas são conferidas estruturalmente, sem exigir recálculo por Excel/LibreOffice.
    Isso não certifica o valor exibido num cache de cálculo antigo.
    """
    column = get_column_letter(coverage_header.index('Localização neste catálogo') + 1)
    requirements = (
        ('Não localizadas neste catálogo', 'não localizado'),
        ('Ambíguas: mais de um registro possível, nenhum escolhido', 'ambíguo'),
    )
    info = {}
    ok = True
    for label, prefix in requirements:
        expected = sum(str(r.get('Localização neste catálogo') or '').startswith(prefix) for r in coverage_rows)
        cells = [r[1] for r in readme_rows if len(r) >= 2 and r[0] == label]
        formula = f'=COUNTIF(\'Cobertura PR53\'!${column}:${column},"{prefix}*")'
        value = cells[0] if len(cells) == 1 else None
        matched = len(cells) == 1 and (value == formula or
                  (isinstance(value, (int, float)) and not isinstance(value, bool) and value == expected))
        info[prefix] = {'esperado': expected, 'ok': matched,
                        'modo': 'fórmula sem recálculo' if value == formula else 'valor'}
        ok = ok and matched
    return ok, info


def check_cht(rows, original_pairs):
    """CHT PR53: os registros em chinês tradicional preservados campo a campo (o PokeData não tem CHT)."""
    cht = [p for p in original_pairs if p.get('Idioma') == 'Chinês tradicional']
    mine = Counter((r['ID EN (PR #53)'],) + tuple(r[b] for _, b in CHT_FIELDS) for r in rows)
    theirs = Counter((p['ID EN'],) + tuple(p[a] for a, _ in CHT_FIELDS) for p in cht)
    return len(rows) == len(cht) and mine == theirs, {'cht': len(rows), 'base_parcial': len(cht)}


def check_comparison(compared, original_pairs):
    """comparacao_pr53.csv: um registro por par confirmado da base parcial; resumo por idioma e resultado."""
    mine = Counter((int(r['id_en']), r['idioma'], r['codigo_pr53']) for r in compared)
    theirs = Counter((p['ID EN'], LANGS.get(p['Idioma'], p['Idioma']), p['Código local completo']) for p in original_pairs)
    info = {'registros': len(compared), 'base_parcial': len(original_pairs),
            'pares_id_idioma': len({(r['id_en'], r['idioma']) for r in compared}),
            'resultado': dict(sorted(Counter(f"{r['idioma']} | {r['resultado']}" for r in compared).items()))}
    return len(compared) == len(original_pairs) and mine == theirs, info


def check_correspondence(cross):
    """Correspondência: toda linha tem ao menos um lado com arte confirmada; linhas = JP + CHS − ambos."""
    jp = sum(r.get('Situação JP') == 'arte confirmada' for r in cross)
    chs = sum(r.get('Situação CN simplificado') == 'arte confirmada' for r in cross)
    both = sum(r.get('Situação JP') == r.get('Situação CN simplificado') == 'arte confirmada' for r in cross)
    info = {'linhas': len(cross), 'jp': jp, 'chs': chs, 'ambos': both,
            'cartas': sum(int(r.get('Cartas do PokeData nesta linha') or 0) for r in cross)}
    return jp + chs - both == len(cross), info


def table(ws, header_row=1):
    rows = ws.iter_rows(min_row=header_row, values_only=True)
    header = [str(h) for h in next(rows)]
    return [dict(zip(header, row)) for row in rows if row and row[0] is not None]


def main(workdir, catalog_path, pr53_path, comparison_path):
    from openpyxl import load_workbook
    manifest_path = Path(workdir) / 'collection_manifest.json'
    manifest = json.loads(manifest_path.read_text(encoding='utf-8')) if manifest_path.exists() else None
    new = load_workbook(catalog_path, read_only=True, data_only=True)
    old = load_workbook(pr53_path, read_only=True, data_only=True)
    formulas = load_workbook(catalog_path, read_only=True, data_only=False)
    try:
        catalogs = {lang: table(new['Catálogo ' + lang]) for lang in ('EN', 'JP', 'CN')}
        original_refs = table(old['Cobertura EN'], 4)
        original_pairs = table(old['Detalhes confirmados'], 4)
        with open(comparison_path, encoding='utf-8-sig', newline='') as stream:
            compared = list(csv.DictReader(stream))
        coverage = table(new['Cobertura PR53'])
        checks = {
            'coleta': check_manifest(manifest),
            'exclusivas': check_exclusives(catalogs, table(new['Exclusivas'])),
            'cobertura_pr53': check_coverage(coverage, original_refs),
            'leia_me_localizacao': check_readme_counts(
                list(formulas['Leia-me'].iter_rows(values_only=True)), coverage,
                next(formulas['Cobertura PR53'].iter_rows(values_only=True))),
            'cht_pr53': check_cht(table(new['CHT PR53']), original_pairs),
            'comparacao_pr53': check_comparison(compared, original_pairs),
            'correspondencia': check_correspondence(table(new['Correspondência'])),
        }
    finally:
        new.close()
        old.close()
        formulas.close()
    report = {name: {'ok': ok, **info} for name, (ok, info) in checks.items()}
    report['catalogo'] = {lang: dict(Counter(r.get('Situação geral') for r in rows if r.get('Carta única') == 1))
                          for lang, rows in catalogs.items()}
    report['escopo'] = 'invariantes e agregados; sem revalidação de imagem nem recálculo de fórmulas'
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if all(ok for ok, _ in checks.values()) else 1


if __name__ == '__main__':
    if len(sys.argv) != 5:
        raise SystemExit(__doc__)
    sys.exit(main(*sys.argv[1:]))
