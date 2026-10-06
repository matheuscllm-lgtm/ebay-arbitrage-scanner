"""Offline audit of the supplied snapshots. Does not validate card identity or prices.

Run from any directory: python research/pokedata/audit_inputs.py
Requires openpyxl. Outputs aggregate metadata only; never exports price values.
"""
import ast
import hashlib
import json
import io
import zipfile
from collections import Counter
from pathlib import Path
from openpyxl import load_workbook

from bridge_partial_ids import build_bridge, id_namespace_check

ROOT = Path(__file__).resolve().parent


def input_bytes(item):
    if "archive" in item:
        with zipfile.ZipFile(ROOT / item["archive"]) as archive:
            return archive.read(item["member"])
    return (ROOT / item["path"]).read_bytes()


def records(workbook, sheet, header_row=1):
    rows = workbook[sheet].iter_rows(min_row=header_row, values_only=True)
    header = next(rows)
    return [dict(zip(header, row)) for row in rows if any(v is not None for v in row)]


def audit():
    report = {"scope": "offline snapshot structure and inherited labels; no image revalidation"}
    manifest = json.loads((ROOT / "inputs/manifest.json").read_text(encoding="utf-8"))
    for item in manifest:
        data = input_bytes(item)
        assert len(data) == item["bytes"], item["path"]
        assert hashlib.sha256(data).hexdigest() == item["sha256"], item["path"]
    report["original_files_sha256_verified"] = len(manifest)
    scripts = list((ROOT / "pipeline").glob("*.py"))
    for path in scripts:
        ast.parse(path.read_text(encoding="utf-8"), filename=str(path))
    report["python_scripts_syntax_checked"] = len(scripts)
    catalog = next(item for item in manifest if item["path"] == "PokeData_catalogo_correspondencia.xlsx")
    w = load_workbook(io.BytesIO(input_bytes(catalog)), read_only=True, data_only=True)
    catalogs = {lang: records(w, "Catálogo " + lang) for lang in ("EN", "JP", "CN")}
    report["catalog"] = {}
    for lang, rows in catalogs.items():
        units = [r for r in rows if r["Carta única"] == 1]
        ids = [r["ID PokeData"] for r in rows]
        report["catalog"][lang] = {"records": len(rows), "units": len(units),
            "inherited_statuses": dict(Counter(r["Situação geral"] for r in units)),
            "duplicate_pokedata_ids": len(ids) - len(set(ids))}
    cor = records(w, "Correspondência")
    report["correspondence"] = {"rows": len(cor),
        "JP": sum(r["Situação JP"] == "confirmado" for r in cor),
        "CHS": sum(r["Situação CN simplificado"] == "confirmado" for r in cor),
        "both": sum(r["Situação JP"] == r["Situação CN simplificado"] == "confirmado" for r in cor),
        "inherited_confidence": dict(Counter(r["Confiança"] for r in cor))}
    report["exclusive_criteria"] = dict(Counter(r["Critério"] for r in records(w, "Exclusivas")))
    sets = records(w, "Sets")
    report["sets"] = dict(Counter(r["Idioma"] for r in sets))
    report["CHT_reference_sets_only"] = len(records(w, "Ref. chinês tradicional"))
    errors = []
    for sheet in w:
        for row in sheet:
            for cell in row:
                if cell.data_type == "e": errors.append(f"{sheet.title}!{cell.coordinate}:{cell.value}")
    report["cached_excel_errors"] = errors
    w.close()
    w = load_workbook(ROOT / "inputs/PokeData_correspondencias_JP_CHS_CHT_parcial_recebido.xlsx", read_only=True, data_only=True)
    en = records(w, "Cobertura EN", 4)
    confirmed = records(w, "Detalhes confirmados", 4)
    pending = records(w, "Inconclusivas", 4)
    ids = {r["ID EN"] for r in en}
    report["partial"] = {"references": len(en), "unique_ids": len(ids),
        "all_raw_above_40": all(isinstance(r["Raw EN (USD)"], (int, float)) and r["Raw EN (USD)"] > 40 for r in en),
        "confirmed_rows": len(confirmed), "confirmed_references": len({r["ID EN"] for r in confirmed}),
        "unique_reference_language_pairs": len({(r["ID EN"], r["Idioma"]) for r in confirmed}),
        "languages": dict(Counter(r["Idioma"] for r in confirmed)),
        "pending_statuses": dict(Counter(r["Status"] for r in pending)),
        "orphan_ids": len({r["ID EN"] for r in confirmed + pending} - ids)}
    # `ID EN` is the partial base's row number (1..3490), not a PokeData ID:
    # integer overlap is coincidence, so bridge by set + number + name instead.
    report["partial"]["id_en_is_row_number"] = sorted(ids) == list(range(1, len(en) + 1))
    report["partial"]["id_namespace"] = id_namespace_check(en, catalogs["EN"])
    report["partial"]["metadata_bridge"] = dict(Counter(b["status"] for b in build_bridge(en, catalogs["EN"])))
    w.close()
    w = load_workbook(ROOT / "pokedata_sets_por_idioma.xlsx", read_only=True, data_only=True)
    report["standalone_sets"] = dict(Counter(r["Idioma"] for r in records(w, "Todos os sets")))
    w.close()
    return report


if __name__ == "__main__":
    print(json.dumps(audit(), ensure_ascii=False, indent=2))
