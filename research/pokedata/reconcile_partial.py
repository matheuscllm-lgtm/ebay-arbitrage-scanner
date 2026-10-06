"""Conciliação offline dos pares confirmados da base parcial com o catálogo de 04/10.

Portado do PR #55 (sessão Claude sharp-planck, `revise_snapshot.py`), só a parte de
conciliação: as reclassificações C1/C2/C4 daquele PR foram superadas pelo pipeline
revisado (commit 0fb962b). Usa a ponte de IDs (`bridge_partial_ids.py`) e as mesmas
normalizações de código do `pipeline/compare_pr53.py` (`pipeline/reference_match.py`).

Compara com o snapshot histórico em inputs/: o catálogo revisado é privado e o
reprocessamento está pendente. Rótulos herdados não são revalidados aqui.

Uso: python research/pokedata/reconcile_partial.py [-o conciliacao.csv]
O CSV (gitignored) traz só identidade e resultado; nunca preços.
"""
import argparse
import csv
import io
import json
import sys
import zipfile
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / "pipeline"))
from bridge_partial_ids import CATALOG_MEMBER, PARTIAL, _records, build_bridge  # noqa: E402
from reference_match import local_code, same_print  # noqa: E402

LANG = {"Japonês": "JP", "Chinês simplificado": "CN", "Chinês tradicional": "CHT"}
SITUATION = {"JP": "Equivalente JP", "CN": "Equivalente CN simplificado"}
RESOLVED = ("unico", "unico_numero_literal")
FIELDS = ["id_en", "carta_en", "idioma", "codigo_local", "id_pokedata", "resultado"]


def catalog_codes(row, lang):
    """Códigos equivalentes de um idioma ('JP: CLL 003; s8a-P 001 | CN: CSV8C 188'), normalizados."""
    parts = dict(p.split(": ", 1) for p in str(row.get("Códigos equivalentes") or "").split(" | ") if ": " in p)
    return {c for c in (local_code(x) for x in parts.get(lang, "").split("; ")) if c}


def reconcile(pair, bridge_row, catalog_by_id):
    """Resultado de um registro confirmado da base parcial frente ao catálogo."""
    lang = LANG[pair["Idioma"]]
    if lang == "CHT":
        return "preservado: CHT é frente própria; o catálogo não tem chinês tradicional"
    if bridge_row["status"] not in RESOLVED:
        return f"sem conciliação: ponte {bridge_row['status']}"
    row = catalog_by_id[int(bridge_row["ids_pokedata"])]
    codes = catalog_codes(row, lang)
    local = local_code(pair["Código local completo"])
    if local is None:
        return "código local ilegível"
    match = same_print(local, codes)
    if match == "igual":
        return "igual"
    if match == "subproduto":
        return "igual no número: subproduto não distinguido no catálogo"
    if codes:
        return "diverge: catálogo aponta outra impressão"
    return f"sem equivalente no catálogo: {row.get(SITUATION[lang])}"


def load():
    from openpyxl import load_workbook
    with zipfile.ZipFile(ROOT / "inputs/pokedata_crossref.zip") as archive:
        data = archive.read(CATALOG_MEMBER)
    w = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    catalog = _records(w, "Catálogo EN", 1)
    w.close()
    w = load_workbook(PARTIAL, read_only=True, data_only=True)
    references, confirmed = _records(w, "Cobertura EN", 4), _records(w, "Detalhes confirmados", 4)
    w.close()
    return catalog, references, confirmed


def run(catalog, references, confirmed):
    bridge = {b["id_en_parcial"]: b for b in build_bridge(references, catalog)}
    by_id = {r["ID PokeData"]: r for r in catalog}
    rows, counts, pairs = [], Counter(), set()
    for pair in confirmed:
        result = reconcile(pair, bridge[pair["ID EN"]], by_id)
        counts[(LANG[pair["Idioma"]], result.split(":")[0])] += 1
        pairs.add((pair["ID EN"], pair["Idioma"]))
        rows.append({"id_en": pair["ID EN"], "carta_en": pair["Carta inglesa — referência"],
                     "idioma": pair["Idioma"], "codigo_local": pair["Código local completo"],
                     "id_pokedata": bridge[pair["ID EN"]]["ids_pokedata"], "resultado": result})
    summary = {"registros": len(confirmed), "pares_id_idioma": len(pairs),
               "resultado": {" | ".join(k): v for k, v in sorted(counts.items())}}
    return rows, summary


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("-o", "--output", help="CSV de auditoria (gitignored); omitido = só resumo")
    args = parser.parse_args(argv)
    rows, summary = run(*load())
    if args.output:
        with open(args.output, "w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=FIELDS)
            writer.writeheader()
            writer.writerows(rows)
    print(json.dumps(summary, ensure_ascii=False, indent=1))
    return 0


if __name__ == "__main__":
    sys.exit(main())
