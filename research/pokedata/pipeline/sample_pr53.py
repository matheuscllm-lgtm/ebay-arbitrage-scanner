"""Apoio — folhas de contato para a revalidação visual pedida em docs/POKEDATA_REPROCESSAMENTO.md, §5.

Uso (na pasta de trabalho): python ../sample_pr53.py <planilha do PR #53> <comparacao_pr53.csv> [n_provaveis] [semente]
Saídas locais (não versionar): revalidacao_pr53.jpg (pares do PR #53 que ficaram "provável"/"inconclusiva"
aqui, lado a lado com a candidata) e revalidacao_provaveis.jpg (amostra aleatória da aba Prováveis).
"""
import csv, pickle, random, sys
from collections import defaultdict
from openpyxl import load_workbook
from common import *
from reference_match import local_code, number_norm, resolve_reference, set_norm
from sheet import sheet

src, comp = sys.argv[1], sys.argv[2]
n_prov = int(sys.argv[3]) if len(sys.argv) > 3 else 48
seed = int(sys.argv[4]) if len(sys.argv) > 4 else 20261006
R = pickle.load(open('result.pkl', 'rb')); CONF, PROV, GRAY = R['CONF'], R['PROV'], R['GRAY']
sets, cards = load(); units = art_units(cards)
LANG = {k: sets[k[0]]['language'] for k in units}
en_by = defaultdict(list)
for k in units:
    if LANG[k] == 'ENGLISH': en_by[(sets[k[0]]['name'].lower(), number_norm(k[1]))].append(k)

def unit_code(u): return (set_norm(sets[u[0]]['code'] or ''), number_norm(u[1]))

# 1) pares do PR #53 que não ficaram "igual": referência EN × candidata desta rodada
wanted = {(r['id_en'], r['idioma'], r['codigo_pr53']): r['resultado'] for r in csv.DictReader(open(comp, encoding='utf-8-sig', newline=''))
          if r['idioma'] in ('JP', 'CHS') and ('provável' in r['resultado'] or 'inconclusiva' in r['resultado'])}
ws = load_workbook(src, read_only=True, data_only=True)['Detalhes confirmados']
hdr = None; pairs = []; labels = []
for row in ws.iter_rows(min_row=4, values_only=True):
    if hdr is None: hdr = [str(h) for h in row]; continue
    if row[0] is None: continue
    d = dict(zip(hdr, row))
    idioma = {'Japonês': 'JP', 'Chinês simplificado': 'CHS'}.get(d['Idioma'])
    key_ = (str(d['ID EN']), idioma, d['Código local completo'])
    if key_ not in wanted: continue
    cands = [dict(id=x['id'], nome=x['name'], unidade=k, numero=k[1]) for k in en_by.get((str(d['Set EN']).lower(), number_norm(d['Código EN'])), []) for x in units[k]]
    uk, _, _ = resolve_reference(d['Carta inglesa — referência'], d['Código EN'], cands)
    loc = local_code(d['Código local completo'])
    if uk is None or loc is None: continue
    t = 'EN_JA' if idioma == 'JP' else 'EN_CH'
    pool = list(PROV[t].get(uk, {})) + list(GRAY[t].get(uk, {}))
    match = [u for u in pool if unit_code(u)[1] == loc[1] and (unit_code(u)[0] == loc[0] or not unit_code(u)[0])] or pool[:1]
    for u in match[:1]:
        pairs.append((uk, u)); labels.append(f"{d['ID EN']} {idioma} {d['Código local completo']}: {wanted[key_][:28]}")
sheet(pairs, 'revalidacao_pr53.jpg', labels, cols=3, h=360)
print('revalidacao_pr53.jpg', len(pairs), 'pares')

# 2) amostra aleatória de pares "provável" (todas as combinações de idiomas)
allp = [(a, b, t) for t in PROV for a in PROV[t] for b in PROV[t][a]]
random.seed(seed); sample = random.sample(allp, min(n_prov, len(allp)))
sheet([(a, b) for a, b, _ in sample], 'revalidacao_provaveis.jpg', [f"{t} {sets[a[0]]['code'] or sets[a[0]]['name'][:12]} {a[1]} × {sets[b[0]]['code'] or sets[b[0]]['name'][:12]} {b[1]}" for a, b, t in sample], cols=4, h=300)
print('revalidacao_provaveis.jpg', len(sample), 'de', len(allp), 'pares prováveis; semente', seed)
