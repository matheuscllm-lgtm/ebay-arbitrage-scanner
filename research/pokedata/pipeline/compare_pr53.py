"""Comparação com a planilha do PR #53 (PokeData_correspondencias_JP_CHS_CHT_parcial.xlsx).

Para cada correspondência "Confirmada" daquela planilha (aba Detalhes confirmados), procura a mesma
carta em inglês neste trabalho e diz se o código local coincide com algum equivalente confirmado aqui.
Uso: python compare_pr53.py <planilha do PR #53> <saida.csv>
Entrada local: result.pkl, catalog.pkl. A saída não contém preços.
"""
import sys, re, csv, pickle
from collections import Counter, defaultdict
from datetime import datetime
from openpyxl import load_workbook
from common import *

src, out = sys.argv[1], sys.argv[2]
rows_, uinfo = pickle.load(open('catalog.pkl', 'rb'))
R = pickle.load(open('result.pkl', 'rb')); CONF, PROV, VIA, STATUS, GRAY = R['CONF'], R['PROV'], R['VIA'], R['STATUS'], R['GRAY']
sets, cards = load(); units = art_units(cards)
LANG = {k: sets[k[0]]['language'] for k in units}
def numnorm(n):
    n = str(n).strip()
    m = re.fullmatch(r'0*(\d+)', n)
    return m.group(1) if m else n.upper()
en_by = defaultdict(list)
for k in units:
    if LANG[k] == 'ENGLISH': en_by[(sets[k[0]]['name'].lower(), numnorm(k[1]))].append(k)
REV_CJ = defaultdict(dict)
for a, c in CONF['CH_JA'].items():
    for b, v in c.items(): REV_CJ[b][a] = v
def mine(k, lang):
    """equivalentes desta rodada (arte confirmada, prováveis, candidatas abaixo do limite), como (código do set, número)"""
    t = 'EN_JA' if lang == 'Japonês' else 'EN_CH'
    conf = list(CONF[t].get(k, {})) + (list(VIA.get(k, {})) if t == 'EN_CH' else [])
    f = lambda u: (setnorm(sets[u[0]]['code'] or ''), numnorm(u[1]), sets[u[0]]['name'])
    return [f(u) for u in conf], [f(u) for u in PROV[t].get(k, {})], [f(u) for u in GRAY[t].get(k, {})]
def setnorm(c):
    """normaliza o código do set: SM4+ = SM4p; promos SV-P; subprodutos 151C1..151C4 = 151C"""
    c = str(c).strip().replace('+', 'p')
    if re.fullmatch(r'151C\d', c, re.I): c = '151C'
    if c.upper().startswith('PROMOSV'): c = 'SV-P'
    return key(c)
def parse_local(code):
    """'M4 114/083' -> ('m4', '114'); 'CBB5C 08 07/07' -> ('cbb5c', '807'); 'PROMOSV08 092/SV-P' -> ('svp', '92')"""
    parts = str(code).strip().split(None, 1)
    if len(parts) < 2: return setnorm(code), ''
    rest = re.sub(r'/[^/\s]*$', '', parts[1])
    digits = re.sub(r'\D', '', rest)
    return setnorm(parts[0]), (str(int(digits)) if digits else rest.upper())
HAS = None
def catalog_state(lang, lc, ln):
    """onde o código do PR #53 cai no catálogo do PokeData usado nesta rodada"""
    L = 'JAPANESE' if lang == 'Japonês' else 'CHINESE'
    sids = [s['id'] for s in sets.values() if s['language'] == L and setnorm(s['code'] or '') == lc]
    if not sids: return 'set fora do catálogo do PokeData'
    us = [k for k in units if k[0] in sids and numnorm(k[1]) == ln]
    if not us: return 'número fora do catálogo do PokeData'
    if all(STATUS[k].get('ENGLISH', '').startswith('inconclusivo: carta sem imagem') for k in us): return 'carta sem imagem no PokeData'
    return 'carta presente com imagem, arte não bateu: revisar'

ws = load_workbook(src, read_only=True, data_only=True)['Detalhes confirmados']
hdr = None; res = []; c = Counter()
for row in ws.iter_rows(min_row=4, values_only=True):
    if hdr is None: hdr = [str(h) for h in row]; continue
    if row[0] is None: continue
    d = dict(zip(hdr, row))
    lang = d['Idioma']; idioma = {'Japonês': 'JP', 'Chinês simplificado': 'CHS', 'Chinês tradicional': 'CHT'}.get(lang, lang)
    base = dict(id_en=d['ID EN'], carta_en=d['Carta inglesa — referência'], codigo_en=d['Código EN'], set_en=d['Set EN'], idioma=idioma, codigo_pr53=d['Código local completo'], nome_local=d['Nome local'])
    if idioma == 'CHT':
        res.append(dict(base, resultado='só no PR #53: chinês tradicional não existe no PokeData', codigos_claude='')); c[(idioma, 'só no PR #53 (CHT)')] += 1; continue
    cands = en_by.get((str(d['Set EN']).lower(), numnorm(d['Código EN'])), [])
    nk = key(split_name(str(d['Carta inglesa — referência']))[0])
    best = [k for k in cands if k[2] == nk] or [k for k in cands if k[2] in nk or nk in k[2]] or cands
    if not best:
        res.append(dict(base, resultado='carta em inglês não localizada nesta rodada', codigos_claude='')); c[(idioma, 'EN não localizada')] += 1; continue
    lc, ln = parse_local(d['Código local completo'])
    conf = []; prov = []; gray = []
    for k in best:
        a, pv, g = mine(k, lang); conf += a; prov += pv; gray += g
    txt = '; '.join(sorted({f"{(x[0] or x[2]).upper() if x[0] else x[2]} {x[1]}" for x in conf}))
    same = lambda xs: any(x[1] == ln and (x[0] == lc or not x[0]) for x in xs)
    if any(x[0] == lc and x[1] == ln for x in conf): r = 'igual (arte confirmada aqui)'
    elif any(x[1] == ln and not x[0] for x in conf): r = 'mesmo número; PokeData não traz o código do set'
    elif same(prov) and conf: r = 'provável aqui, com a mesma candidata'
    elif conf: r = 'outra impressão confirmada aqui; a do PR #53: ' + catalog_state(lang, lc, ln)
    elif same(prov): r = 'provável aqui, com a mesma candidata'
    elif same(gray): r = 'inconclusiva aqui, com a mesma candidata'
    else: r = 'sem confirmação aqui: ' + catalog_state(lang, lc, ln)
    res.append(dict(base, resultado=r, codigos_claude=txt)); c[(idioma, r)] += 1
with open(out, 'w', newline='', encoding='utf-8-sig') as f:
    w = csv.DictWriter(f, fieldnames=['id_en', 'carta_en', 'codigo_en', 'set_en', 'idioma', 'codigo_pr53', 'nome_local', 'resultado', 'codigos_claude'])
    w.writeheader(); w.writerows(res)
print('registros confirmados no PR #53:', len(res))
for k in sorted(c): print(f'  {k[0]:4} {c[k]:4}  {k[1]}')
