"""Revisão offline do snapshot com as regras corrigidas (C1, C2, C4) e conciliação dos pares.

O catálogo publicado não pode ser regenerado sem os intermediários do pipeline
(`*.pkl`, `ext/`, `img/`). Este script aplica ao snapshot as mesmas funções de
`pipeline/identity.py` que o pipeline corrigido usa e mede o efeito de cada
correção. Os snapshots ficam intactos; a saída é um registro de auditoria.

Uso: python research/pokedata/revise_snapshot.py [-o revisao.csv]
O CSV (gitignored) traz só identidade e status; nunca preços.
"""
import argparse
import csv
import io
import re
import sys
import zipfile
from collections import Counter
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(ROOT / "pipeline"))
from identity import CANDIDATE, match_level, number_key, single_language  # noqa: E402
from bridge_partial_ids import (CATALOG_MEMBER, PARTIAL, _records, build_bridge,  # noqa: E402
                                text_key)

LANG_OF = {"Inglês": "ENGLISH", "Japonês": "JAPANESE", "Chinês simplificado": "CHINESE"}
PARTIAL_LANG = {"Japonês": "JP", "Chinês simplificado": "CN", "Chinês tradicional": "CHT"}
CATALOG_LANGS = ("ENGLISH", "JAPANESE", "CHINESE")
FIELDS = ["frente", "idioma", "referencia", "status_anterior", "status_novo", "motivo", "metodo", "data"]


def exclusive_evidence(lang, criterion):
    """Evidência externa citada no critério herdado da aba Exclusivas."""
    if lang == "ENGLISH" and "Limitless" in criterion:
        return {"JAPANESE": "ausente"}
    if lang == "JAPANESE" and "Limitless também não lista" in criterion:
        return {"ENGLISH": "ausente"}
    return {}


def reclassify_exclusive(lang, criterion):
    statuses = {t: "não encontrado" for t in CATALOG_LANGS if t != lang}
    return single_language(lang, statuses, exclusive_evidence(lang, criterion))


def old_sig(row):
    """Chave antiga da aba Correspondência: set + primeiro inteiro + 1º parceiro JP/CN."""
    first = re.search(r"(\d+)", str(row["Número"]))
    codes = dict(part.split(": ", 1) for part in str(row["Códigos equivalentes"] or "").split(" | ") if ": " in part)
    partners = tuple(codes.get(lang, "").split("; ")[0] for lang in ("JP", "CN"))
    return (row["set_id"], int(first.group(1)) if first else 10 ** 6) + partners


def new_sig(row):
    """Chave corrigida: número completo + nome base + 1º parceiro JP/CN."""
    return (row["set_id"], number_key(row["Número"]), text_key(row["Nome base"])) + old_sig(row)[2:]


def local_code(code):
    """Código local → (set, número) comparável entre a base parcial e o catálogo.

    'M4 114/083' → ('M4', '114'); 'SM4p 120/114' e 'SM4+ 120' → ('SM4P', '120');
    promo 'PROMOSV08 092/SV-P' → ('SV-P', '92'), como no catálogo;
    pacote de gemas 'CBB5C 08 07/07' → ('CBB5C', '807'), como 'CBB5C 0807' no catálogo.
    """
    parts = str(code or "").strip().split()
    if len(parts) < 2:
        return None
    set_code, number = parts[0], " ".join(parts[1:])
    head, _, denominator = number.partition("/")
    if re.fullmatch(r"[A-Za-z]+-P", denominator.strip()):
        set_code = denominator.strip()
    elif len(parts) == 3 and "/" in parts[2]:
        head = parts[1] + parts[2].split("/")[0]
    return set_code.upper().replace("+", "P"), number_key(head)


def same_print(local, codes):
    """'igual', 'subproduto' (mesmo número; set do catálogo é prefixo, ex. 151C4 × 151C) ou None."""
    if local in codes:
        return "igual"
    if any(n == local[1] and local[0].startswith(s) and local[0] != s for s, n in codes):
        return "subproduto"
    return None


def catalog_codes(row, lang):
    """Códigos equivalentes de um idioma na linha do Catálogo EN, normalizados."""
    codes = dict(part.split(": ", 1) for part in str(row["Códigos equivalentes"] or "").split(" | ") if ": " in part)
    return {local_code(c) for c in codes.get(lang, "").split("; ") if local_code(c)}


def reconcile(pair, bridge_row, catalog_by_id):
    """Compara um par confirmado da base parcial com o catálogo, via ponte de metadados."""
    lang = PARTIAL_LANG[pair["Idioma"]]
    if lang == "CHT":
        return "preservado: CHT é frente própria; catálogo não tem chinês tradicional"
    if bridge_row["status"] != "unico":
        return f"sem conciliação: ponte {bridge_row['status']}"
    row = catalog_by_id[int(bridge_row["ids_pokedata"])]
    codes = catalog_codes(row, lang)
    match = same_print(local_code(pair["Código local completo"]), codes)
    if match == "igual":
        return "igual"
    if match:
        return "igual no número: subproduto não distinguido no catálogo"
    if codes:
        return "diverge: catálogo aponta outra impressão"
    situation = row["Equivalente JP" if lang == "JP" else "Equivalente CN simplificado"]
    return f"sem equivalente no catálogo: {situation}"


def load():
    from openpyxl import load_workbook
    with zipfile.ZipFile(ROOT / "inputs/pokedata_crossref.zip") as archive:
        data = archive.read(CATALOG_MEMBER)
    w = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    out = {"catalog_en": _records(w, "Catálogo EN", 1), "exclusive": _records(w, "Exclusivas", 1),
           "correspondence": _records(w, "Correspondência", 1)}
    w.close()
    w = load_workbook(PARTIAL, read_only=True, data_only=True)
    out["references"] = _records(w, "Cobertura EN", 4)
    out["confirmed"] = _records(w, "Detalhes confirmados", 4)
    w.close()
    return out


def revise(data):
    today = date.today().isoformat()
    log, summary = [], {}

    # C1 — exclusivas: só com ausência comprovada em todos os idiomas-alvo
    c1 = Counter()
    for row in data["exclusive"]:
        lang = LANG_OF[row["Idioma"]]
        status, reason = reclassify_exclusive(lang, row["Critério"])
        c1[(row["Idioma"], status, "com fonte externa" if exclusive_evidence(lang, row["Critério"]) else "só ausência")] += 1
        log.append({"frente": "C1 exclusividade", "idioma": row["Idioma"],
                    "referencia": f"{row['Set']} {row['Número']} {row['Nome']}",
                    "status_anterior": "exclusiva", "status_novo": status, "motivo": reason,
                    "metodo": "identity.single_language sobre o critério herdado", "data": today})
    summary["C1_exclusivas"] = {" | ".join(k): v for k, v in sorted(c1.items())}

    # C2 — deduplicação da aba Correspondência
    units = [r for r in data["catalog_en"] if r["Carta única"] == 1 and r["Situação geral"] == "confirmado"]
    old = {old_sig(r) for r in units}
    new = {new_sig(r) for r in units}
    summary["C2_correspondencia"] = {"unidades_EN_confirmadas": len(units), "linhas_chave_antiga": len(old),
                                     "linhas_chave_corrigida": len(new), "linhas_recuperadas": len(new) - len(old),
                                     "linhas_publicadas": len(data["correspondence"])}

    # C4 — arte × impressão
    c4 = Counter()
    for row in data["correspondence"]:
        c4[match_level([row["Pontos coincidentes JP"], row["Pontos coincidentes CN"]])] += 1
    summary["C4_arte_impressao"] = {f"arte {a} / impressão {i}": v for (a, i), v in sorted(c4.items())}

    # Conciliação dos pares confirmados da base parcial, via ponte de metadados (PR #54)
    bridge = {b["id_en_parcial"]: b for b in build_bridge(data["references"], data["catalog_en"])}
    by_id = {r["ID PokeData"]: r for r in data["catalog_en"]}
    rec, pairs = Counter(), set()
    for pair in data["confirmed"]:
        result = reconcile(pair, bridge[pair["ID EN"]], by_id)
        rec[(PARTIAL_LANG[pair["Idioma"]], result.split(":")[0])] += 1
        pairs.add((pair["ID EN"], pair["Idioma"]))
        log.append({"frente": "conciliação base parcial", "idioma": pair["Idioma"],
                    "referencia": f"ID EN {pair['ID EN']} {pair['Carta inglesa — referência']} → {pair['Código local completo']}",
                    "status_anterior": pair["Status"], "status_novo": pair["Status"], "motivo": result,
                    "metodo": "ponte set+número+nome; código local × Códigos equivalentes", "data": today})
    summary["conciliacao"] = {"registros": len(data["confirmed"]), "pares_id_idioma": len(pairs),
                              "resultado": {" | ".join(k): v for k, v in sorted(rec.items())}}
    return log, summary


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("-o", "--output", help="CSV de auditoria (gitignored); omitido = só resumo")
    args = parser.parse_args(argv)
    log, summary = revise(load())
    if args.output:
        with open(args.output, "w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=FIELDS)
            writer.writeheader()
            writer.writerows(log)
    import json
    print(json.dumps(summary, ensure_ascii=False, indent=1))
    return 0


if __name__ == "__main__":
    sys.exit(main())
